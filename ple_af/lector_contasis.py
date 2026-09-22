"""Lectura del Excel "FORMATO 7.1" exportado desde Contasis."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from .utilidades import a_texto, normalizar

# Columna estándar de Contasis para cada dato (respaldo si no se detecta por título)
COLUMNAS_ESTANDAR = {
    "codigo": "B", "cuenta": "C", "descripcion": "D", "marca": "E", "modelo": "F",
    "serie": "G", "saldo_inicial": "H", "adquisiciones": "I", "mejoras": "J",
    "retiros": "K", "otros_ajustes": "L", "valor_historico": "M", "ajuste_inflacion": "N",
    "valor_ajustado": "O", "fecha_adquisicion": "P", "fecha_uso": "Q", "metodo": "R",
    "doc_autorizacion": "S", "porcentaje": "T", "dep_acum_anterior": "U",
    "dep_ejercicio": "V", "dep_retiros": "W", "dep_otros": "X", "dep_acum_historica": "Y",
    "ajuste_inflacion_dep": "Z", "dep_acum_ajustada": "AA",
}

CAMPOS_IMPORTE = [
    "saldo_inicial", "adquisiciones", "mejoras", "retiros", "otros_ajustes",
    "valor_historico", "ajuste_inflacion", "valor_ajustado", "dep_acum_anterior",
    "dep_ejercicio", "dep_retiros", "dep_otros", "dep_acum_historica",
    "ajuste_inflacion_dep", "dep_acum_ajustada",
]

NOMBRES_CAMPOS = {
    "codigo": "Código", "cuenta": "Cuenta contable", "descripcion": "Descripción",
    "marca": "Marca", "modelo": "Modelo", "serie": "Serie/placa",
    "saldo_inicial": "Saldo inicial", "adquisiciones": "Adquisiciones",
    "mejoras": "Mejoras", "retiros": "Retiros/bajas", "otros_ajustes": "Otros ajustes",
    "valor_historico": "Valor histórico", "ajuste_inflacion": "Ajuste por inflación",
    "valor_ajustado": "Valor ajustado", "fecha_adquisicion": "Fecha de adquisición",
    "fecha_uso": "Fecha inicio de uso", "metodo": "Método", "doc_autorizacion": "Doc. autorización",
    "porcentaje": "% depreciación", "dep_acum_anterior": "Dep. acumulada anterior",
    "dep_ejercicio": "Dep. del ejercicio", "dep_retiros": "Dep. retiros/bajas",
    "dep_otros": "Dep. otros ajustes", "dep_acum_historica": "Dep. acumulada histórica",
    "ajuste_inflacion_dep": "Ajuste inflación dep.", "dep_acum_ajustada": "Dep. acumulada ajustada",
}


def _clasificar_titulo(titulo: str) -> str | None:
    """Devuelve la clave del campo según el texto del título (normalizado)."""
    t = normalizar(titulo)
    if not t:
        return None
    reglas = [
        ("codigo", lambda s: s.startswith("CODIGO RELACIONADO")),
        ("cuenta", lambda s: s.startswith("CUENTA CONTABLE")),
        ("descripcion", lambda s: s == "DESCRIPCION"),
        ("marca", lambda s: s.startswith("MARCA")),
        ("modelo", lambda s: s.startswith("MODELO")),
        ("serie", lambda s: "SERIE" in s),
        ("saldo_inicial", lambda s: s.startswith("SALDO INICIAL")),
        ("adquisiciones", lambda s: s.startswith("ADQUISICIONES")),
        ("mejoras", lambda s: s.startswith("MEJORAS")),
        ("retiros", lambda s: s.startswith("RETIROS")),
        ("otros_ajustes", lambda s: s.startswith("OTROS AJUSTE")),
        ("valor_historico", lambda s: s.startswith("VALOR HISTORICO")),
        ("ajuste_inflacion_dep", lambda s: s.startswith("AJUSTE POR INFLACION") and "DEPRECIACION" in s),
        ("ajuste_inflacion", lambda s: s.startswith("AJUSTE POR INFLACION")),
        ("valor_ajustado", lambda s: s.startswith("VALOR AJUSTADO")),
        ("fecha_adquisicion", lambda s: s.startswith("FECHA DE ADQUISICION")),
        ("fecha_uso", lambda s: s.startswith("FECHA DE INICIO")),
        ("metodo", lambda s: s.startswith("METODO")),
        ("doc_autorizacion", lambda s: "AUTORIZACION" in s),
        ("porcentaje", lambda s: s.startswith("PORCENTAJE")),
        ("dep_acum_anterior", lambda s: s.startswith("DEPRECIACION ACUMULADA AL CIERRE")),
        ("dep_retiros", lambda s: s.startswith("DEPRECIACION DEL EJERCICIO RELACIONAD")),
        ("dep_ejercicio", lambda s: s == "DEPRECIACION DEL EJERCICIO"),
        ("dep_otros", lambda s: s.startswith("DEPRECIACION RELACIONAD")),
        ("dep_acum_historica", lambda s: s.startswith("DEPRECIACION ACUMULADA HISTORICA")),
        ("dep_acum_ajustada", lambda s: s.startswith("DEPRECIACION ACUMULADA AJUSTADA")),
    ]
    for clave, condicion in reglas:
        if condicion(t):
            return clave
    return None


@dataclass
class ArchivoContasis:
    ruc: str = ""
    razon_social: str = ""
    ejercicio: int | None = None
    activos: list[dict] = field(default_factory=list)   # cada uno con 'fila' y campos crudos
    totales: dict = field(default_factory=dict)          # fila "Totales:" del Excel
    fila_totales: int | None = None
    columnas: dict = field(default_factory=dict)         # campo -> letra
    avisos_lectura: list[str] = field(default_factory=list)
    filas_sin_codigo: list[int] = field(default_factory=list)


def _leer_filas(contenido: bytes, nombre: str) -> list[list]:
    """Devuelve todas las filas de la hoja como lista de listas (índice 0 = columna A)."""
    if nombre.lower().endswith(".xls"):
        import xlrd
        libro = xlrd.open_workbook(file_contents=contenido)
        hoja = next((h for h in libro.sheets() if "7.1" in h.name), libro.sheet_by_index(0))
        filas = []
        for r in range(hoja.nrows):
            fila = []
            for c in range(hoja.ncols):
                celda = hoja.cell(r, c)
                valor = celda.value
                if celda.ctype == xlrd.XL_CELL_EMPTY or valor == "":
                    valor = None
                elif celda.ctype == xlrd.XL_CELL_DATE:
                    valor = xlrd.xldate.xldate_as_datetime(valor, libro.datemode)
                fila.append(valor)
            filas.append(fila)
        return filas
    libro = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    hoja = next((h for h in libro.worksheets if "7.1" in h.title), libro.worksheets[0])
    filas = [list(f) for f in hoja.iter_rows(values_only=True)]
    libro.close()
    return filas


def leer_excel_contasis(contenido: bytes, nombre: str = "archivo.xlsx") -> ArchivoContasis:
    filas = _leer_filas(contenido, nombre)
    res = ArchivoContasis()

    def celda(r: int, c: int):
        if r < len(filas) and c < len(filas[r]):
            return filas[r][c]
        return None

    # --- Cabecera (EJERCICIO / RUC / RAZÓN SOCIAL) ---
    for r in range(min(15, len(filas))):
        for valor in filas[r]:
            if not isinstance(valor, str):
                continue
            t = normalizar(valor)
            if t.startswith("EJERCICIO") and res.ejercicio is None:
                m = re.search(r"(\d{4})", t)
                if m:
                    res.ejercicio = int(m.group(1))
            elif t.startswith("R.U.C") or t.startswith("RUC"):
                m = re.search(r"(\d{11})", t)
                if m:
                    res.ruc = m.group(1)
            elif t.startswith("RAZON SOCIAL"):
                res.razon_social = valor.split(":", 1)[-1].strip()

    # --- Fila de títulos ---
    fila_titulos = None
    for r in range(min(40, len(filas))):
        if any(_clasificar_titulo(v) == "codigo" for v in filas[r] if isinstance(v, str)):
            fila_titulos = r
            break
    if fila_titulos is None:
        raise ValueError(
            "No se encontró la fila de títulos ('CÓDIGO RELACIONADO CON EL ACTIVO FIJO'). "
            "¿Es la exportación FORMATO 7.1 de Contasis?"
        )

    # --- Columnas por título (en las 4 filas de títulos) ---
    columnas: dict[str, int] = {}
    for r in range(fila_titulos, min(fila_titulos + 4, len(filas))):
        for c, valor in enumerate(filas[r]):
            if isinstance(valor, str):
                clave = _clasificar_titulo(valor)
                if clave and clave not in columnas:
                    columnas[clave] = c
    for clave, letra in COLUMNAS_ESTANDAR.items():
        if clave not in columnas:
            columnas[clave] = column_index_from_string(letra) - 1
            res.avisos_lectura.append(
                f"No se encontró el título de '{NOMBRES_CAMPOS[clave]}'; se usó la columna estándar {letra}."
            )
    res.columnas = {k: get_column_letter(v + 1) for k, v in columnas.items()}

    # --- Datos ---
    col_codigo = columnas["codigo"]
    for r in range(fila_titulos + 4, len(filas)):
        fila = filas[r]
        if not any(v not in (None, "") for v in fila):
            continue
        textos = " ".join(normalizar(v) for v in fila if isinstance(v, str))
        if "TOTALES" in textos:
            res.fila_totales = r + 1
            res.totales = {k: celda(r, c) for k, c in columnas.items() if k in CAMPOS_IMPORTE}
            break
        activo = {k: celda(r, c) for k, c in columnas.items()}
        activo["fila"] = r + 1  # número de fila en Excel
        if activo["codigo"] in (None, ""):
            res.filas_sin_codigo.append(r + 1)
            continue
        activo["codigo_numerico"] = isinstance(activo["codigo"], (int, float))
        activo["codigo"] = a_texto(activo["codigo"])
        res.activos.append(activo)

    if not res.activos:
        res.avisos_lectura.append("El archivo no tiene activos.")
    return res
