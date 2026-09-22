"""Orquesta la generación completa: 7.1, 7.3, 7.4, controles, ZIP y reporte."""
from __future__ import annotations

import io
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config as C
from .archivos import armar_zip, contenido_txt, nombre_archivo
from .config import Config
from .diario import IndiceDiario
from .controles import comparar_anio_anterior, comparar_totales
from .formato_71 import TITULOS_71, Registro71, construir_71
from .formato_73 import TITULOS_73, DatoME, construir_73
from .lector_contasis import ArchivoContasis
from .observaciones import Observaciones
from .utilidades import ruc_valido


@dataclass
class Resultado:
    registros_71: list[Registro71]
    lineas_73: list[list[str]]
    observaciones: Observaciones
    tabla_totales: list[dict]
    archivos: dict[str, bytes] = field(default_factory=dict)   # nombre -> contenido TXT

    @property
    def zip(self) -> bytes:
        return armar_zip(self.archivos)


def validar_config(cfg: Config) -> Observaciones:
    obs = Observaciones()
    if not ruc_valido(cfg.ruc):
        obs.error(f"RUC '{cfg.ruc}' no válido (11 dígitos y dígito verificador).", campo="RUC")
    if cfg.ejercicio is None or cfg.ejercicio < C.EJERCICIO_MINIMO:
        obs.error(f"Ejercicio {cfg.ejercicio}: la app solo trabaja ejercicios {C.EJERCICIO_MINIMO} en adelante.",
                  campo="Ejercicio")
    return obs


def generar(
    archivo: ArchivoContasis,
    cfg: Config,
    datos_me: list[DatoME] | None = None,
    tc_cierre=None,
    archivo_anterior: ArchivoContasis | None = None,
    diario: IndiceDiario | None = None,
) -> Resultado:
    obs = validar_config(cfg)
    if obs.hay_errores:
        return Resultado([], [], obs, [])

    for aviso in archivo.avisos_lectura:
        obs.advertencia(aviso)
    if diario is None:
        obs.advertencia(
            "No se subió el Libro Diario: el CUO y el correlativo (campos 2 y 3) se generaron como "
            f"'{cfg.prefijo_cuo}1', '{cfg.prefijo_asiento}1'… Sube el diario PLE para usar los asientos reales.",
            campo="CUO")
    else:
        for aviso in diario.avisos:
            obs.advertencia(aviso, campo="Libro Diario")
        if diario.ruc and diario.ruc != cfg.ruc:
            obs.error(f"El Libro Diario es de otro RUC ({diario.ruc}).", campo="Libro Diario")
        if diario.ejercicio and diario.ejercicio != cfg.ejercicio:
            obs.error(f"El Libro Diario es del ejercicio {diario.ejercicio}, no de {cfg.ejercicio}.",
                      campo="Libro Diario")
    registros, obs71 = construir_71(archivo, cfg, diario)
    obs.extender(obs71)

    lineas_73: list[list[str]] = []
    if datos_me:
        lineas_73, obs73 = construir_73(registros, datos_me, tc_cierre, cfg)
        obs.extender(obs73)

    tabla, obs_tot = comparar_totales(archivo, registros)
    obs.extender(obs_tot)

    if archivo_anterior is not None:
        if archivo_anterior.ruc and archivo_anterior.ruc != cfg.ruc:
            obs.advertencia(f"El archivo del año anterior es de otro RUC ({archivo_anterior.ruc}); no se comparó.")
        else:
            ej_ant = archivo_anterior.ejercicio or cfg.ejercicio - 1
            cfg_ant = Config(ruc=cfg.ruc, ejercicio=ej_ant)
            registros_ant, _ = construir_71(archivo_anterior, cfg_ant)
            obs.extender(comparar_anio_anterior(registros, registros_ant, ej_ant))

    res = Resultado(registros, lineas_73, obs, tabla)
    if not obs.hay_errores:
        lineas_71 = [r.campos for r in registros]
        res.archivos = {
            nombre_archivo(cfg, C.LIBRO_71, bool(lineas_71)): contenido_txt(lineas_71),
            nombre_archivo(cfg, C.LIBRO_73, bool(lineas_73)): contenido_txt(lineas_73),
            nombre_archivo(cfg, C.LIBRO_74, False): b"",
        }
    return res


def reporte_excel(res: Resultado, cfg: Config) -> bytes:
    """Excel de revisión con 7.1, 7.3, observaciones y totales."""
    wb = Workbook()
    negrita = Font(bold=True)
    relleno = PatternFill("solid", fgColor="DDEBF7")

    def hoja(ws, titulos, filas):
        ws.append(titulos)
        for c in ws[1]:
            c.font = negrita
            c.fill = relleno
        for f in filas:
            ws.append(f)
        for i, t in enumerate(titulos, start=1):
            largo = max([len(str(t))] + [len(str(f[i - 1])) for f in filas if i - 1 < len(f) and f[i - 1] is not None])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, largo + 2), 45)
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "7.1"
    hoja(ws, ["Fila Excel"] + TITULOS_71, [[r.fila] + r.campos for r in res.registros_71])
    hoja(wb.create_sheet("7.3"), TITULOS_73, res.lineas_73)
    obs = res.observaciones.como_filas()
    titulos_obs = ["Tipo", "Fila Excel", "Columna", "Código", "Campo", "Detalle"]
    hoja(wb.create_sheet("Observaciones"), titulos_obs, [[o[t] for t in titulos_obs] for o in obs])
    titulos_tot = ["Concepto", "Calculado", "Excel (fila Totales)", "Diferencia"]
    hoja(wb.create_sheet("Totales"), titulos_tot, [[t[k] for k in titulos_tot] for t in res.tabla_totales])
    info = wb.create_sheet("Datos")
    for fila in (("RUC", cfg.ruc), ("Razón social", cfg.razon_social), ("Ejercicio", cfg.ejercicio),
                 ("Archivos", ", ".join(res.archivos) or "No generados (hay errores)")):
        info.append(fila)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
